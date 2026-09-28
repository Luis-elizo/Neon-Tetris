import docx

doc = docx.Document(r'C:\Users\luisa\Downloads\Tarea2_AngelyVC_NahomyZM_LuisEH..docx')
section = doc.sections[0]
print(f'Margins: top={section.top_margin.inches}, bottom={section.bottom_margin.inches}, left={section.left_margin.inches}, right={section.right_margin.inches}')

for i in range(len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.text.strip():
        fonts = set()
        sizes = set()
        bolds = set()
        colors = set()
        for r in p.runs:
            fonts.add(r.font.name)
            sizes.add(r.font.size.pt if r.font.size else None)
            bolds.add(r.bold)
            colors.add(r.font.color.rgb if r.font.color else None)
        align = p.alignment
        print(f'P{i} [align={align}, style={p.style.name}]: text="{p.text[:80]}"')
        print(f'   runs: fonts={fonts}, sizes={sizes}, bolds={bolds}')
