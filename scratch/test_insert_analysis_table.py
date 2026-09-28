import os
import subprocess
import pypdf
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=20, bottom=20, left=35, right=35):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def update_document(docx_path, out_docx, out_pdf):
    doc = docx.Document(docx_path)
    
    # 1. Encontrar párrafo de análisis de resultados
    target_idx = None
    p_sec5 = None
    p_q3 = None
    
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if "Análisis de resultados:" in txt or "Anlisis de resultados:" in txt:
            target_idx = i
        elif "5. Respuestas a las preguntas de análisis" in txt:
            p_sec5 = p
        elif "Pregunta 3: Necesidad de la lista doblemente enlazada" in txt:
            p_q3 = p

    p_target = doc.paragraphs[target_idx]
    
    # Lead-in paragraph para la tabla de análisis
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.space_before = Pt(4)
    p_intro.paragraph_format.space_after = Pt(2)
    p_intro.paragraph_format.keep_with_next = True
    p_intro.paragraph_format.widow_control = True
    r_intro = p_intro.add_run("Análisis e Interpretación de Resultados:")
    r_intro.font.name = 'Arial'
    r_intro.font.size = Pt(11)
    r_intro.bold = True
    
    # Crear la tabla de análisis
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    encabezados = ["Caso / Tamaño (N)", "Resultado Observado", "Interpretación y Causa Técnica"]
    anchos = [Inches(1.2), Inches(1.8), Inches(3.8)]
    
    hdr = tbl.rows[0]
    set_table_header(hdr)
    prevent_row_split(hdr)
    for idx, nombre in enumerate(encabezados):
        cell = hdr.cells[idx]
        cell.width = anchos[idx]
        set_cell_background(cell, "EAEAEA")
        set_cell_margins(cell, top=20, bottom=20, left=30, right=30)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(nombre)
        r.font.name = 'Arial'
        r.font.size = Pt(8.5)
        r.bold = True
        
    filas_analisis = [
        (
            "N = 10",
            "Empate técnico\n(0.90 µs vs 1.00 µs)",
            "Bubble Sort es una fracción más rápido porque opera directamente sobre el arreglo sin el sobrecosto de llamadas recursivas ni reserva de memoria para la mezcla."
        ),
        (
            "N = 100",
            "Merge Sort 3.05x más rápido\n(38.50 µs frente a 117.50 µs)",
            "Se supera el umbral de eficiencia donde el enfoque recursivo divide y vencerás sobrepasa a las comparaciones cuadráticas adyacentes."
        ),
        (
            "N = 1 000",
            "Merge Sort 18.97x más rápido\n(0.65 ms frente a 12.33 ms)",
            "La brecha es clara en la fluidez del juego: Bubble Sort toma más de 12 ms, mientras que Merge Sort resuelve el ordenamiento de forma instantánea."
        ),
        (
            "N = 10 000",
            "Merge Sort 223.51x más rápido\n(5.86 ms vs 1 309.15 ms)",
            "Bubble Sort supera los 1.3 segundos (provocaría congelamiento de pantalla perceptible), mientras Merge Sort procesa todo en menos de 6 ms a 60 FPS estables."
        ),
        (
            "Crecimiento\n(1k → 10k)",
            "Bubble aumentó 106x\nMerge aumentó solo 9x",
            "Confirma la teoría: al multiplicar datos por 10, Bubble Sort escala de forma cuadrática O(n²), mientras Merge Sort muestra un crecimiento casi lineal O(n log n)."
        )
    ]
    
    for f_idx, fila in enumerate(filas_analisis):
        row = tbl.add_row()
        prevent_row_split(row)
        color = "F9F9F9" if f_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(fila):
            cell = row.cells[c_idx]
            cell.width = anchos[c_idx]
            set_cell_background(cell, color)
            set_cell_margins(cell, top=14, bottom=14, left=25, right=25)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1)
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(8.0)
            if c_idx == 0 or c_idx == 1:
                r.bold = True

    # Reubicar los elementos XML en la posición de p_target:
    parent = p_target._p.getparent()
    p_target_idx = parent.index(p_target._p)
    
    parent.insert(p_target_idx, p_intro._p)
    parent.insert(p_target_idx + 1, tbl._tbl)
    parent.remove(p_target._p)

    # Asegurar salto de página ANTES de Sección 5 para que no quede huérfana en Página 4
    if p_sec5 is not None:
        p_br_sec5 = doc.add_paragraph()
        r_br = p_br_sec5.add_run()
        r_br.add_break(docx.enum.text.WD_BREAK.PAGE)
        p_sec5_parent = p_sec5._p.getparent()
        p_sec5_parent.insert(p_sec5_parent.index(p_sec5._p), p_br_sec5._p)

    # Asegurar salto de página ANTES de Pregunta 3 para distribución equitativa (Página 6)
    if p_q3 is not None:
        p_br_q3 = doc.add_paragraph()
        r_br3 = p_br_q3.add_run()
        r_br3.add_break(docx.enum.text.WD_BREAK.PAGE)
        p_q3_parent = p_q3._p.getparent()
        p_q3_parent.insert(p_q3_parent.index(p_q3._p), p_br_q3._p)

    doc.save(out_docx)
    print(f"Documento guardado: {out_docx}")

if __name__ == '__main__':
    src = r"c:\Users\luisa\Universidad\ll ciclo 2026\Estructura de Datos\Neon-Tetris\INFORME_PROYECTO1_EIF207.docx"
    out_docx = r"c:\Users\luisa\Universidad\ll ciclo 2026\Estructura de Datos\Neon-Tetris\INFORME_PROYECTO1_EIF207.docx"
    out_pdf = r"c:\Users\luisa\Universidad\ll ciclo 2026\Estructura de Datos\Neon-Tetris\INFORME_PROYECTO1_EIF207.pdf"
    
    update_document(src, out_docx, out_pdf)

    ps_script = f'''
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    try {{
        $doc = $word.Documents.Open("{out_docx}", $false, $true)
        $doc.SaveAs([ref]"{out_pdf}", [ref]17)
        $doc.Close()
        Write-Host "Exportacion exitosa de PDF"
    }} finally {{
        $word.Quit()
    }}
    '''
    ps_path = r"c:\Users\luisa\Universidad\ll ciclo 2026\Estructura de Datos\Neon-Tetris\scratch\export_final.ps1"
    open(ps_path, 'w', encoding='utf-8').write(ps_script)
    subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', ps_path])

    r = pypdf.PdfReader(out_pdf)
    print('Total pages in final PDF:', len(r.pages))
    for i, page in enumerate(r.pages):
        txt = page.extract_text().strip()
        first_line = txt.split('\n')[0] if txt else ''
        last_line = txt.split('\n')[-1] if txt else ''
        print(f'=== PAGE {i+1} (chars: {len(txt)}) ===')
        print('  First:', first_line[:85])
        print('  Last: ', last_line[:85])
