import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def generar_docx_normal(ruta_grafico, ruta_docx):
    doc = docx.Document()

    # Configuración de página: 1 pulgada en todos los márgenes (estilo estándar)
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Estilo base
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)

    # Helper para párrafos normales con interlineado 1.5 y espaciado normal
    def agregar_p(texto="", bold_prefix="", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0, line_spacing=1.5):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)

        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = 'Times New Roman'
            r_b.font.size = Pt(12)
            r_b.bold = True

        if texto:
            r_t = p.add_run(texto)
            r_t.font.name = 'Times New Roman'
            r_t.font.size = Pt(12)

        return p

    def agregar_titulo_seccion(texto):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(texto)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(16)
        r.bold = True
        return p

    def agregar_subtitulo(texto):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(texto)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        return p

    # -------------------------------------------------------------
    # PORTADA (Igual al modelo de Tarea 2 del estudiante)
    # -------------------------------------------------------------
    for _ in range(2):
        agregar_p("", space_after=12)

    agregar_p("Sede Regional Brunca", align=WD_ALIGN_PARAGRAPH.CENTER, bold_prefix="", space_after=4)
    doc.paragraphs[-1].runs[0].bold = True

    agregar_p("Campus Pérez Zeledón", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    for _ in range(2):
        agregar_p("", space_after=12)

    agregar_p("Curso: EIF207 Estructuras de Datos", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    agregar_p("Proyecto 1: Neon Tetris", align=WD_ALIGN_PARAGRAPH.CENTER, bold_prefix="", space_after=24)
    doc.paragraphs[-1].runs[0].bold = True

    for _ in range(2):
        agregar_p("", space_after=12)

    agregar_p("Estudiante:", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    agregar_p("Luis Andres Elizondo Hernandez", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    for _ in range(2):
        agregar_p("", space_after=12)

    agregar_p("Profesor: [Nombre del profesor]", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    agregar_p("II Ciclo 2026", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    for _ in range(2):
        agregar_p("", space_after=12)

    agregar_p("27 de septiembre del 2026", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    # Salto de página tras la portada
    doc.add_page_break()

    # -------------------------------------------------------------
    # INTRODUCCIÓN
    # -------------------------------------------------------------
    agregar_titulo_seccion("Introducción")

    agregar_p(
        "En este primer proyecto del curso de Estructuras de Datos desarrollamos una versión funcional del juego clásico Tetris, "
        "llamada \"Neon Tetris\", programada en C++ y utilizando la biblioteca gráfica Raylib. El propósito principal del proyecto "
        "fue aplicar de forma práctica las estructuras de datos lineales que estudiamos en las primeras semanas del ciclo (pilas, colas, "
        "listas enlazadas, colas de prioridad y algoritmos de ordenamiento), pero con la condición indispensable de implementarlas "
        "desde cero mediante nodos, punteros y memoria dinámica, sin usar ninguna de las estructuras de la biblioteca estándar (STL)."
    )

    agregar_p(
        "Cada una de las mecánicas que tiene el juego está conectada directamente con una estructura de datos específica. Por ejemplo, "
        "la fila de piezas que van a salir se maneja con una cola; la mecánica de guardar una pieza para después (hold) se resuelve con una pila; "
        "el historial que permite deshacer jugadas y ver el replay usa una lista doblemente enlazada; los eventos aleatorios que modifican la "
        "partida usan una cola de prioridad basada en un montículo (heap); y el tablero del juego se representa como una lista enlazada de filas. "
        "Además, implementamos dos métodos de ordenamiento (Burbuja y Merge Sort) para organizar la tabla de mejores puntajes y comparamos sus "
        "tiempos de ejecución reales para ver si coincidían con la teoría de complejidad asintótica."
    )

    agregar_p(
        "En este informe explicamos la arquitectura y la relación entre todas estas estructuras, detallamos la complejidad de cada operación "
        "implementada, presentamos las pruebas de rendimiento con sus gráficos y respondemos a las preguntas de análisis requeridas."
    )

    # -------------------------------------------------------------
    # DESARROLLO
    # -------------------------------------------------------------
    agregar_titulo_seccion("Desarrollo")

    agregar_subtitulo("1. Descripción del juego y arquitectura general")
    agregar_p(
        "El juego sigue la lógica básica de Tetris en un tablero de 10 columnas por 20 filas. Las piezas caen de forma constante y el jugador "
        "puede moverlas hacia los lados, rotarlas y acelerar su caída. Cuando se completa una línea horizontal, esta se limpia, sumando puntos "
        "y haciendo que los bloques superiores desciendan. Si los bloques acumulados llegan arriba y una nueva pieza no puede entrar, la partida "
        "termina (Game Over)."
    )

    agregar_p(
        "Para mantener el código ordenado y fácil de mantener, organizamos el proyecto en tres carpetas principales:\n"
        "• estructuras: Contiene las clases de las estructuras de datos hechas a mano con sus nodos y punteros (Cola, Pila, ListaDoble, "
        "ColaPrioridad, Tablero y ManejadorJSON).\n"
        "• logica: Contiene la definición de las piezas con sus rotaciones (Pieza) y las funciones de ordenamiento (ordenamientos).\n"
        "• ui: Maneja las distintas pantallas del juego (Menú principal, Pantalla de juego, Pausa, Game Over y Ranking) junto con los colores neón."
    )

    agregar_subtitulo("2. Relación y funcionamiento de las estructuras de datos")
    agregar_p(
        "En lugar de ver las estructuras como elementos aislados, cada una cumple un rol esencial dentro del ciclo de juego en main.cpp:"
    )

    agregar_p(
        "En Tetris las piezas no deben salir completamente al azar porque podrían repetirse muchas veces la misma o pasar mucho tiempo "
        "sin que salga una pieza larga (la I). Por eso se usa el sistema de \"bolsa de 7\": se toman las 7 piezas clásicas, se mezclan aleatoriamente "
        "con el algoritmo de Fisher-Yates y se meten a la cola. Conforme el jugador va jugando, se van sacando con desencolar() y en pantalla se muestran "
        "las siguientes 3 piezas mirando los primeros nodos de la cola. Cuando la cola se vacía, se genera otra bolsa de 7 y se vuelve a llenar.",
        bold_prefix="• Cola de piezas futuras (Cola): "
    )

    agregar_p(
        "La mecánica de \"Hold\" permite guardar la pieza actual para usarla en otro momento. Esto lo modelamos con una pila de capacidad 1. "
        "Si la pila está vacía y el jugador presiona Shift, la pieza actual se mete a la pila y se saca una nueva de la cola de piezas. Si ya había una pieza "
        "guardada, se intercambian: se desapila la que estaba guardada para pasarla al juego y se apila la que estaba cayendo. Para que el jugador no abuse "
        "de esto, solo se permite un cambio por turno hasta que la pieza caiga y se fije en el tablero.",
        bold_prefix="• Pila para la pieza en espera (Pila): "
    )

    agregar_p(
        "El tablero de 20x10 no es un simple arreglo estático bidimensional, sino una estructura de filas y celdas enlazadas mediante punteros en "
        "cuatro direcciones (arriba, abajo, izquierda, derecha). Cuando se llena una fila, no hay que copiar todos los números de arriba para bajarlos: "
        "simplemente se desenganchan los punteros de la fila llena, se liberan sus nodos con delete, se conectan los vecinos y se crea una fila vacía "
        "nueva arriba en el tope.",
        bold_prefix="• Tablero como lista de filas (Tablero): "
    )

    agregar_p(
        "Cada vez que una pieza se coloca en el tablero o se limpian líneas, se guarda una copia del tablero y del puntaje en un nodo de la lista doble. "
        "Como cada nodo tiene punteros siguiente y anterior, durante la partida el jugador puede presionar la tecla Z para deshacer jugadas (retroceder en el "
        "historial) o la tecla Y para rehacerlas (avanzar en el historial). Cuando la partida termina en Game Over, esta misma lista sirve para el modo Replay, "
        "donde se puede volver a ver toda la partida paso a paso o en reproducción automática.",
        bold_prefix="• Lista doblemente enlazada para historial y replay (ListaDoble): "
    )

    agregar_p(
        "Para hacer el juego más dinámico, programamos eventos que ocurren cada cierto tiempo: un aumento temporal de velocidad (\"Fiebre de velocidad\"), "
        "una línea de basura que sube desde el fondo (\"Terremoto\") o un bono de +500 puntos. Estos eventos se guardan en una cola de prioridad implementada "
        "como un montículo mínimo (Min-Heap), donde la prioridad es la marca de tiempo (timestamp) en la que deben ocurrir. De esta forma, el evento que tiene "
        "que ejecutarse más pronto siempre queda en la raíz (al frente), y cuando el tiempo de la partida llega a esa marca, se saca y se aplica su efecto.",
        bold_prefix="• Cola de prioridad para eventos programados (ColaPrioridad): "
    )

    agregar_p(
        "Los mejores puntajes se guardan en un archivo scores.json. Desde el menú de Ranking, el jugador puede ver la lista de récords ordenada "
        "de mayor a menor y puede elegir si quiere ordenarla usando Bubble Sort o Merge Sort, viendo en pantalla cuántos microsegundos tardó cada algoritmo.",
        bold_prefix="• Persistencia y ordenamiento de puntajes: "
    )

    agregar_subtitulo("3. Análisis de complejidad teórica (Notación Big-O)")
    agregar_p(
        "En la siguiente tabla resumimos la complejidad teórica de las operaciones principales de las estructuras implementadas en el proyecto:"
    )

    # TABLA COMPLEJIDAD (Formato limpio, Times New Roman, encabezado gris claro)
    tbl_comp = doc.add_table(rows=1, cols=7)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_comp.autofit = False

    encabezados_c = ["Estructura", "Operación", "Mejor", "Promedio", "Peor", "Espacio", "Descripción breve"]
    anchos_c = [Inches(1.1), Inches(1.1), Inches(0.6), Inches(0.7), Inches(0.7), Inches(0.6), Inches(1.7)]

    hdr = tbl_comp.rows[0]
    for idx, nombre in enumerate(encabezados_c):
        cell = hdr.cells[idx]
        cell.width = anchos_c[idx]
        set_cell_background(cell, "EAEAEA")
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(nombre)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.bold = True

    filas_c = [
        ("Cola FIFO", "encolar()", "O(1)", "O(1)", "O(1)", "O(1)", "Inserta directo al final usando el puntero ultimo."),
        ("Cola FIFO", "desencolar()", "O(1)", "O(1)", "O(1)", "O(1)", "Saca del frente y libera la memoria con delete."),
        ("Cola FIFO", "obtenerEn(k)", "O(1)", "O(k)", "O(n)", "O(1)", "Recorre hasta la posición k para ver piezas futuras."),
        ("Pila LIFO", "apilar() / desapilar()", "O(1)", "O(1)", "O(1)", "O(1)", "Modifica directamente el nodo que está en el tope."),
        ("Lista Doble", "agregarEstado()", "O(1)", "O(1)", "O(1)", "O(F*C)", "Agrega al final; copia la matriz de 20x10."),
        ("Lista Doble", "retroceder() [Undo]", "O(1)", "O(1)", "O(1)", "O(1)", "Mueve el puntero: actual = actual->anterior."),
        ("Lista Doble", "avanzar() [Redo]", "O(1)", "O(1)", "O(1)", "O(1)", "Mueve el puntero: actual = actual->siguiente."),
        ("Cola Prioridad", "encolar()", "O(1)", "O(log n)", "O(log n)", "O(1)", "Inserta al final y flota hacia arriba (sift-up)."),
        ("Cola Prioridad", "desencolar()", "O(1)", "O(log n)", "O(log n)", "O(1)", "Saca la raíz y reacomoda hacia abajo (sift-down)."),
        ("Cola Prioridad", "verTiempoFrente()", "O(1)", "O(1)", "O(1)", "O(1)", "Accede directo a la raíz arreglo[0]."),
        ("Tablero", "limpiarLineas()", "O(F*C)", "O(F*C)", "O(F*C)", "O(1)", "Reconecta punteros de filas llenas y crea nuevas arriba."),
        ("Ordenamiento", "Bubble Sort", "O(n)", "O(n²)", "O(n²)", "O(1)", "Compara e intercambia elementos vecinos."),
        ("Ordenamiento", "Merge Sort", "O(n log n)", "O(n log n)", "O(n log n)", "O(n)", "Divide recursivamente y mezcla sub-arreglos.")
    ]

    for f_idx, fila in enumerate(filas_c):
        row = tbl_comp.add_row()
        color = "F9F9F9" if f_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(fila):
            cell = row.cells[c_idx]
            cell.width = anchos_c[c_idx]
            set_cell_background(cell, color)
            set_cell_margins(cell, top=40, bottom=40, left=50, right=50)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            if c_idx in [2, 3, 4, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.0)
            if c_idx in [2, 3, 4, 5]:
                r.bold = True

    agregar_p("", space_after=6)

    agregar_subtitulo("4. Comparación empírica y pruebas de rendimiento")
    agregar_p(
        "Para comparar el rendimiento real de Bubble Sort (O(n²)) contra Merge Sort (O(n log n)), realizamos un benchmark con la librería "
        "<chrono> de C++, midiendo el tiempo de ejecución en microsegundos (µs). Probamos con arreglos aleatorios de 10, 100, 1 000 y 10 000 "
        "registros con nombres sintéticos y puntajes al azar, promediando 5 repeticiones por cada tamaño para asegurar datos representativos."
    )

    # TABLA BENCHMARK
    tbl_bench = doc.add_table(rows=1, cols=6)
    tbl_bench.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_bench.autofit = False

    encabezados_b = ["Tamaño (N)", "Bubble Sort (µs)", "Bubble Sort (ms)", "Merge Sort (µs)", "Merge Sort (ms)", "Diferencia"]
    anchos_b = [Inches(1.0), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1)]

    hdr_b = tbl_bench.rows[0]
    for idx, nombre in enumerate(encabezados_b):
        cell = hdr_b.cells[idx]
        cell.width = anchos_b[idx]
        set_cell_background(cell, "EAEAEA")
        set_cell_margins(cell, top=50, bottom=50, left=50, right=50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(nombre)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.bold = True

    datos_b = [
        ("N = 10", "0.90 µs", "0.0009 ms", "1.00 µs", "0.0010 ms", "Empate técnico (0.9x)"),
        ("N = 100", "117.50 µs", "0.1175 ms", "38.50 µs", "0.0385 ms", "Merge 3x más rápido"),
        ("N = 1 000", "12 333.00 µs", "12.33 ms", "650.00 µs", "0.65 ms", "Merge 19x más rápido"),
        ("N = 10 000", "1 309 150.00 µs", "1 309.15 ms", "5 857.00 µs", "5.86 ms", "Merge 223.5x más rápido")
    ]

    for f_idx, fila in enumerate(datos_b):
        row = tbl_bench.add_row()
        color = "F9F9F9" if f_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(fila):
            cell = row.cells[c_idx]
            cell.width = anchos_b[c_idx]
            set_cell_background(cell, color)
            set_cell_margins(cell, top=40, bottom=40, left=50, right=50)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.0)
            if c_idx == 5:
                r.bold = True

    agregar_p("", space_after=6)

    # Gráfico embebido
    if os.path.exists(ruta_grafico):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(ruta_grafico, width=Inches(6.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run("Figura 1: Gráfico comparativo de tiempos entre Bubble Sort y Merge Sort en escala logarítmica.")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(9.5)
        r_cap.italic = True

    agregar_p(
        "Al analizar los números obtenidos notamos varios puntos importantes:\n"
        "• Con N = 10, los dos algoritmos tardan prácticamente lo mismo (0.9 µs vs 1.0 µs). De hecho, Bubble Sort fue una fracción "
        "más rápido porque no tiene que hacer llamadas recursivas ni reservar memoria extra para el arreglo temporal de mezcla.\n"
        "• Con N = 100, Merge Sort ya empieza a ganar ventaja y es 3 veces más rápido (38.5 µs frente a 117.5 µs).\n"
        "• Con N = 1 000, la diferencia ya es notable: Merge Sort tarda solo 0.65 ms mientras que Bubble Sort tarda 12.33 ms (unas 19 veces más rápido).\n"
        "• Con N = 10 000, la diferencia es enorme: Bubble Sort tarda más de 1.3 segundos (1 309.15 ms), lo que en un juego causaría que la pantalla "
        "se congele de forma molesta, mientras que Merge Sort hace todo el trabajo en menos de 6 milisegundos (5.86 ms), siendo más de 223 veces más rápido.\n"
        "• Si nos fijamos en cómo crecieron los tiempos de 1 000 a 10 000 (multiplicando los datos por 10), Bubble Sort aumentó su tiempo unas 106 veces "
        "(12.33 ms a 1 309.15 ms), lo cual coincide casi exacto con lo que predice la teoría cuadrática (10² = 100). En cambio, Merge Sort solo aumentó "
        "9 veces (0.65 ms a 5.86 ms), confirmando su comportamiento casi lineal O(n log n)."
    )

    agregar_subtitulo("5. Respuestas a las preguntas obligatorias de la sección 7")

    # Pregunta 1
    agregar_p(
        "Pregunta 1: Compare el tiempo de ordenar la tabla de puntajes con el algoritmo de complejidad O(n²) contra el de O(n log n), "
        "para tamaños crecientes de datos de prueba. ¿A partir de qué tamaño se nota la diferencia? ¿Coincide con lo que predice la notación asintótica?",
        bold_prefix=""
    )
    doc.paragraphs[-1].runs[0].bold = True

    agregar_p(
        "Con pocos datos (N = 10), los dos algoritmos duran casi lo mismo (0.9 µs y 1.0 µs), porque para arreglos pequeños el costo de las llamadas "
        "recursivas y la memoria dinámica que usa Merge Sort empata el tiempo con las pocas comparaciones directas que hace Bubble Sort. "
        "La diferencia se empieza a notar a partir de N = 100, donde Merge Sort ya es 3 veces más veloz. Ya con N = 1 000 la diferencia es clara (19x), "
        "y en N = 10 000 la diferencia es gigantesca, pues Bubble tarda más de 1.3 segundos y Merge menos de 6 milisegundos (223.5x más rápido).\n"
        "Sí coincide totalmente con lo que dice la teoría asintótica: Bubble Sort tiene una complejidad cuadrática O(n²), por lo que al multiplicar "
        "los datos por 10 su tiempo se multiplica por 100 (en nuestras pruebas reales subió 106 veces). Mientras tanto, Merge Sort es O(n log n) y crece "
        "de forma mucho más suave, siendo la opción adecuada para manejar grandes cantidades de datos."
    )

    # Pregunta 2
    agregar_p(
        "Pregunta 2: Explique por qué representar el tablero como una lista enlazada de filas es una forma razonable de modelar la limpieza de líneas, "
        "y qué costo (en operaciones) tiene en su implementación insertar una fila vacía y eliminar una fila completa.",
        bold_prefix=""
    )
    doc.paragraphs[-1].runs[0].bold = True

    agregar_p(
        "En una matriz común de toda la vida (un arreglo estático de 20x10), cuando una fila se llena en el medio, para hacer que los bloques de arriba "
        "\"caigan\" hay que hacer ciclos que copien fila por fila hacia abajo, moviendo decenas o cientos de números en la memoria. En cambio, al modelar "
        "el tablero como una lista enlazada de filas, las filas de arriba no se tienen que mover de lugar en la memoria. Lo único que se hace es soltar "
        "los punteros de arriba y abajo que estaban conectados a la fila llena, y conectar directamente la fila de arriba con la de abajo. La fila completa "
        "se borra con delete y se inserta una fila vacía nueva en la parte superior conectando sus punteros.\n"
        "• Costo de eliminar una fila completa: Como el tablero tiene 10 columnas, se recorren las 10 celdas desconectando sus enlaces verticales y "
        "liberando la memoria. Esto toma exactamente 10 liberaciones y unas 20 reasignaciones de punteros, lo que representa un costo de O(C) operaciones "
        "(tiempo constante O(1) porque las 10 columnas nunca cambian).\n"
        "• Costo de insertar una fila vacía: Se crean 10 nodos celda nuevos con new, se enlazan horizontalmente y se conectan sus punteros hacia abajo con "
        "la fila que antes estaba de primera. Esto toma 10 creaciones y unas 30 asignaciones de puntero, teniendo un costo de O(C) operaciones (O(1) constante)."
    )

    # Pregunta 3
    agregar_p(
        "Pregunta 3: ¿Por qué el replay requiere una lista doblemente enlazada y no una simplemente enlazada? ¿Qué costo tiene, en su implementación, "
        "retroceder o avanzar un paso?",
        bold_prefix=""
    )
    doc.paragraphs[-1].runs[0].bold = True

    agregar_p(
        "Para poder hacer un replay o para deshacer y rehacer jugadas, necesitamos movernos hacia adelante y hacia atrás. Si usáramos una lista simplemente "
        "enlazada, cada nodo solo sabe quién es el siguiente. Para avanzar en el tiempo no hay problema (actual = actual->siguiente), pero para retroceder "
        "una sola jugada no se puede regresar directamente porque no hay puntero al anterior. Habría que empezar a buscar desde el puro inicio de la lista "
        "recorriendo nodo por nodo hasta encontrar cuál era el que apuntaba a la jugada actual. Esto costaría O(k) operaciones cada vez que el usuario presione "
        "la tecla de retroceder (donde k es el número de la jugada actual), lo que causaría tirones o bajones de FPS si la partida ya lleva muchos turnos. "
        "La lista doblemente enlazada soluciona esto porque cada nodo tiene un puntero anterior y un puntero siguiente.\n"
        "• Costo de avanzar un paso: Solo se hace actual = actual->siguiente. Es una sola asignación de puntero, costo O(1).\n"
        "• Costo de retroceder un paso: Solo se hace actual = actual->anterior. Es una sola asignación de puntero, costo O(1)."
    )

    # Pregunta 4
    agregar_p(
        "Pregunta 4: ¿Cómo garantiza su implementación que la cola de eventos programados siempre mantenga al frente el evento más próximo a dispararse? "
        "¿Qué complejidad tiene insertar un nuevo evento?",
        bold_prefix=""
    )
    doc.paragraphs[-1].runs[0].bold = True

    agregar_p(
        "Nuestra clase ColaPrioridad está implementada con un Montículo Binario Mínimo (Min-Heap) sobre un arreglo dinámico. La propiedad que siempre "
        "se cumple en este montículo es que cualquier nodo padre siempre tiene un tiempo menor o igual al de sus dos nodos hijos. Gracias a esta regla "
        "matemática, el evento que tiene el tiempo más pequeño (el que va a pasar más pronto) siempre queda garantizado en la raíz del árbol, que es la "
        "posición 0 del arreglo. Saber cuál es el siguiente evento con verTiempoFrente() cuesta O(1) porque solo se lee esa primera casilla.\n"
        "• Inserción de un nuevo evento (encolar): El nuevo evento se coloca al final del arreglo y luego se compara con su padre (sift-up o flotación). "
        "Si su tiempo es menor que el del padre, se intercambian y se sigue subiendo hasta que quede en su lugar correcto. Como un árbol binario con n elementos "
        "tiene una altura de log₂ n, a lo sumo se hacen log₂ n comparaciones. Por lo tanto, la inserción tiene una complejidad de O(log n), lo cual es muy "
        "eficiente y cumple la regla de no insertar al final y reordenar todo con un ordenamiento general (que costaría O(n log n))."
    )

    # -------------------------------------------------------------
    # CONCLUSIONES
    # -------------------------------------------------------------
    agregar_titulo_seccion("Conclusiones")

    agregar_p(
        "Desarrollar este proyecto sin utilizar la STL nos permitió entender a fondo cómo funcionan las estructuras de datos lineales por dentro, "
        "cómo se manejan los punteros y la importancia de liberar siempre la memoria dinámica con delete para evitar fugas.",
        bold_prefix="• "
    )

    agregar_p(
        "Cada mecánica del juego encaja de forma natural con una estructura específica: la cola para las piezas futuras, la pila para el hold, "
        "la lista de filas para el tablero, la lista doble para el historial y el replay, y el montículo para los eventos por tiempo.",
        bold_prefix="• "
    )

    agregar_p(
        "Las pruebas de rendimiento demostraron en la práctica lo que se estudia en la teoría: para pocos datos algoritmos simples como Bubble Sort "
        "funcionan bien y son fáciles de programar, pero conforme el volumen de datos crece, algoritmos eficientes como Merge Sort son indispensables "
        "para que una aplicación no se vuelva lenta o se congele.",
        bold_prefix="• "
    )

    doc.save(ruta_docx)
    print(f"Documento normal generado en: {ruta_docx}")

if __name__ == '__main__':
    base = r"c:\Users\luisa\Universidad\ll ciclo 2026\Estructura de Datos\Neon-Tetris"
    graf = os.path.join(base, "scratch", "benchmark_chart.png")
    docx_path = os.path.join(base, "INFORME_PROYECTO1_EIF207.docx")
    generar_docx_normal(graf, docx_path)
